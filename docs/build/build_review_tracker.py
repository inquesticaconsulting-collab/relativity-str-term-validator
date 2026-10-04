#!/usr/bin/env python3
"""Build docs/STR_Validation_Rules_Team_Review.xlsx and docs/RULES.md.

The rule catalog below is the single source for both outputs.
Run from the repository root:  python docs/build/build_review_tracker.py
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DOCS = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

HEAD_BG = "3D5A80"
SEV_FILL = {"ERROR": "F4CCCC", "WARNING": "FFF2CC", "INFO": "D9E7F5"}
FILLIN = "FFF9DB"          # pale yellow = cells the team fills in
EXAMPLE_BG = "EFEFEF"
THIN = Side(style="thin", color="C0C7D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def F(**kw):
    kw.setdefault("size", 10)
    return Font(name="Arial", **kw)


HEADF = Font(name="Arial", size=10, bold=True, color="FFFFFF")

# (code, severity, what it detects, handling, example, discussion point)
RULES = [
    ("OVER_450_CHARS", "ERROR", "Term exceeds the STR 450-character limit.",
     "Relativity rejects it. Not auto-fixable (NEEDS REVIEW): split into several lines — each line is its own query and reports separately.", "—",
     "Any matters where we expect very long negotiated strings? Agree who does the splitting — analyst or requestor."),
    ("UNBALANCED_QUOTES", "ERROR", "Odd number of double quotes; phrase never closed.",
     "Query is invalid/unpredictable. Auto-balanced — verify the revision.",
     '"unclosed phrase  →  "unclosed phrase"',
     "OK to trust the auto-balance, or always send back to requestor to confirm the intended phrase?"),
    ("UNBALANCED_PARENS", "ERROR", "Parentheses do not balance.",
     "Grouping is ambiguous or invalid. Auto-balanced — verify the revision.",
     "(missing close  →  (missing close)",
     "Same question as quotes — accept auto-repair or confirm with requestor?"),
    ("EMPTY_PARENS", "ERROR", "Empty group ().",
     "Nothing to search; removed automatically.", "term ()  →  term",
     "Likely uncontroversial — adopt as-is?"),
    ("COMMA_SEPARATED_LIST", "ERROR", "Several terms pasted on one line, separated by commas/semicolons (outside quotes).",
     "dtSearch treats commas as spaces — the line runs as ONE exact phrase and hits almost nothing. Rewritten with OR; splitting to separate lines is better for per-term counts.",
     "apple, pear  →  apple OR pear",
     "Default fix: split to separate lines (per-term counts) or join with OR (one count)? Pick a house rule."),
    ("STRAY_SEPARATOR", "INFO", "A leading/trailing comma or semicolon.",
     "Carries no meaning in dtSearch; removed.", "apple,  →  apple",
     "Adopt as-is?"),
    ("LEADING/TRAILING/DOUBLE_OPERATOR", "ERROR", "Operator with a missing operand at start/end, or two binary operators in a row.",
     "Invalid query. Auto-repaired by dropping the dangling operator; in 'AND w/10' the proximity operator is kept.",
     "contract AND AND breach  →  contract AND breach",
     "A dangling operator can mean a term was lost in copy/paste. Always ask requestor, or only when it looks suspicious?"),
    ("OPERATOR_AFTER/BEFORE_PAREN", "ERROR", "Binary operator immediately inside a group boundary.",
     "A group cannot begin or end with an operator; auto-repaired.", "( AND x)  →  (x)",
     "Adopt as-is?"),
    ("NOT_BEFORE_OPERATOR", "ERROR", "NOT directly followed by AND/OR ('NOT OR').",
     "Invalid. The NOT is dropped and the operator kept — but this pattern often means AND NOT (exclusion) was intended: the two readings are opposites.",
     "x NOT OR y  →  x OR y (confirm AND NOT wasn't meant)",
     "Always route this one back to the requestor?"),
    ("PROX_SPACING", "ERROR", "Space inside a proximity operator.",
     "dtSearch only recognizes w/N or pre/N with no spaces; otherwise 'w' is searched literally. Auto-corrected.",
     "invoice w/ 5 payment  →  invoice W/5 payment",
     "Adopt as-is?"),
    ("PROX_TRAILING_SLASH", "ERROR", "Proximity with a trailing slash, e.g. w/7/.",
     "Unrecognizable to dtSearch (the terms would run as a phrase). Auto-corrected to w/7.",
     "a w/7/ b  →  a W/7 b",
     "Adopt as-is?"),
    ("PROX_DISTANCE_ASSUMED (was PROX_MALFORMED)", "WARNING", "w/ or pre/ with no number.",
     "A default of w/5 is inserted so the term runs, labeled ASSUMPTION APPLIED for confirmation.",
     "invoice w/ payment  →  invoice W/5 payment (assumed)",
     "Is w/5 the right house default? Confirm or change the assumed distance."),
    ("PROX_ZERO", "WARNING", "Distance of zero (w/0).",
     "Auto-replaced with w/1 (within one word, either order), labeled ASSUMPTION APPLIED. A quoted phrase is the ordered alternative.",
     "apple w/0 pear  →  apple W/1 pear (assumed)",
     "Accept w/1 as the default, or prefer converting to a quoted phrase?"),
    ("LEADING_WILDCARD", "ERROR", "Wildcard at the start of a word.",
     "Prohibited in STRs and a documented performance hazard. Removed; use trailing wildcard, stemming (~), or explicit OR variants.",
     "*fraud*  →  fraud*",
     "Meaning-changing. Standard substitute: stemming, or enumerate variants with OR? Who drafts the variants?"),
    ("FUZZY_NOT_SUPPORTED", "ERROR", "Fuzzy operator %.",
     "Not supported in STRs. Removed; pre-expand via the dtSearch Dictionary and join variants with OR.",
     "app%le  →  apple (plus Dictionary variants)",
     "Who runs the Dictionary expansion — analyst or admin? Add it to the checklist?"),
    ("STEM_LEADING", "ERROR", "~ at the start of a word.",
     "Stemming goes at the END of the root (English only). Moved there automatically.",
     "~apply  →  apply~",
     "Adopt as-is?"),
    ("STRAY_TILDE", "ERROR", "~~ or ~~~ attached to a word or quoted phrase (not a numeric range).",
     "Invalid decoration — removed. ~~ is only valid between integers (12~~24); a single ~ only as stemming on a bare word.",
     '"confidential"~~  →  "confidential"',
     "Adopt as-is?"),
    ("PHONIC_MISPLACED", "WARNING", "# inside or after a word.",
     "Removed as a likely typo (labeled ASSUMPTION APPLIED); # is only valid immediately before a word.",
     "proj#ect  →  project (assumed typo)",
     "Agree: remove-as-typo is the default; ask the requestor only when phonic seems plausible?"),
    ("NUMERIC_RANGE_MALFORMED", "ERROR", "Invalid N~~M numeric range.",
     "Both bounds must be integers. The intended range can't be guessed, so this stays NEEDS REVIEW.",
     "12~~abc  →  (human fixes, e.g. 12~~24)",
     "Adopt as-is?"),
    ("WILDCARD_ONLY", "ERROR", "Term is only wildcards.",
     "Matches every word in the index (including xfirstword/xlastword) — returns all documents.",
     "*  →  (replace with a real term)",
     "Adopt as-is (strike and query requestor)?"),
    ("NOISE_WORD_TERM", "ERROR", "Standalone term is a dtSearch noise word.",
     "Returns zero hits on a default index. NEEDS REVIEW: replace, or have the admin remove it from the noise list and rebuild.",
     "the  →  (strike or replace)",
     "Confirm our indexes use the default noise list. When would we ever rebuild instead of striking?"),
    ("FIELD_SYNTAX_UNSUPPORTED", "ERROR", "Metadata written as field:value (custodian:, subject:, date:, ext:, folder:, project:).",
     "Not dtSearch syntax — an STR term searches content only. Extracted to the 'Extracted Conditions' sheet; they belong in the searchable-set saved search. Values are validated (impossible dates, misspelled extensions and field names).",
     "custodian:Smith AND breach  →  breach (condition listed for the saved search)",
     "Agree the house rule: metadata goes in the saved search. Who applies the extracted conditions? Re-add subject: phrases as content phrases?"),
    ("UNSUPPORTED_FUNCTION", "ERROR", "Function-style syntax dtSearch doesn't have, e.g. stemming().",
     "Only date(), mail(), creditcard() exist. Removed; for stemming, use word~.",
     "stemming()  →  (removed)",
     "Adopt as-is?"),
    ("MIXED_AND_OR", "WARNING", "AND and OR mixed with no parentheses.",
     "dtSearch applies OR before AND. The revision adds parentheses showing how the query actually runs — confirm intent.",
     "a AND b OR c  →  a AND (b OR c)",
     "Always send the explicit-parentheses version back to the requestor for confirmation, or only on request?"),
    ("AND_INSIDE_PROX_GROUP", "WARNING", "Group joined by w/N or pre/N contains AND.",
     "Documented ambiguous pattern with unpredictable results. Use OR inside proximity-connected groups, or restructure.",
     "(a AND b) w/10 (c AND d)  →  (a w/10 c) AND (b w/10 d)",
     "The restructure changes meaning — treat as meaning-changing (needs requestor approval)?"),
    ("NESTED_PROX / MANY_PROX_OPS", "WARNING", "Chained/nested or numerous proximity operators.",
     "Discouraged for performance; chained proximity can evaluate ambiguously. Split across lines.",
     "a w/3 b w/3 c  →  (a w/3 b) AND (b w/3 c)",
     "Current threshold: warn at 3+ proximity operators. Right level?"),
    ("PROX_TOO_WIDE", "WARNING", "Very wide proximity window (above ~25 words).",
     "Behaves like AND across paragraphs; inflates hits. Confirm width.",
     "a w/50 b  →  confirm or narrow",
     "Threshold currently 25 words — agree the number as a team standard."),
    ("NOISE_WORD_IN_QUERY", "WARNING", "Noise word inside an unquoted expression.",
     "The word is skipped by the index; in phrases it matches any single word in that position. Verify counts make sense.",
     "statement of work — 'of' acts as a placeholder",
     "Should the memo to requestors routinely explain the placeholder effect, or only when counts look off?"),
    ("RESERVED_WORD_UNQUOTED", "WARNING", "'to' or 'contains' used as an ordinary word outside quotes.",
     "These are dtSearch connector words; quote the phrase so they are searched literally.",
     'go to market  →  "go to market"',
     "Adopt as-is?"),
    ("SPECIAL_CHAR", "WARNING", "Character outside the default searchable alphabet (&, _, :, @, etc.).",
     "Not matched literally by a default index (AT&T indexes as 'at' + 't'). Verify alphabet file; consider phrase forms, mail(), or regex.",
     'AT&T  →  "at t" (default index)',
     "Need admin to document our alphabet-file settings per workspace. Who owns getting that list?"),
    ("HYPHENATED_TERM", "WARNING", "Hyphenated word.",
     "Hyphen handling is index-configurable (commonly behaves as a space). Search all variants.",
     'e-mail  →  ("e-mail" OR email OR "e mail")',
     "Auto-propose the 3-variant OR form as our default? It broadens scope — approval needed each time?"),
    ("UNICODE_CHARS", "WARNING", "Smart quotes/dashes/non-breaking spaces from Word.",
     "Not treated as phrase delimiters or indexed as typed. Auto-replaced with ASCII equivalents.",
     "“phrase”  →  \"phrase\"",
     "Adopt as-is (syntax-only fix, notify not approve)?"),
    ("EARLY_WILDCARD / MULTI_WILDCARD", "WARNING", "* near word start, or several * in one word.",
     "Performance and over-inclusion risk. Narrow or enumerate variants.",
     "b*n*a  →  banana OR ...",
     "How aggressive should we be here — advise only, or push back on requestors?"),
    ("MULTI_WILDCARD_RUN", "WARNING", "Consecutive asterisks (**).",
     "Behaves like a single * but signals a typo. Auto-collapsed to one *.",
     ".doc**  →  .doc*",
     "Adopt as-is?"),
    ("X_WILDCARD_TRAP", "WARNING", "Pattern matching built-in xfirstword/xlastword tokens.",
     "Those tokens exist in every document — the pattern returns everything.",
     "x*  →  narrow the pattern",
     "Adopt as-is?"),
    ("OVERBROAD_TERM", "WARNING", "Very short standalone term.",
     "Matches enormous volume. Add context via proximity or qualifiers.",
     "ab  →  ab w/10 <context>",
     "Threshold currently ≤3 characters. Exempt a list of known acronyms per matter?"),
    ("REDUNDANT_OPERAND", "WARNING", "Same word on both sides of AND/OR.",
     "Equivalent to the single word; usually a paste slip. Collapsed — check whether a different term was intended.",
     "deal OR deal  →  deal",
     "Always query the requestor (a different second term was probably intended)?"),
    ("DUPLICATE_TERM", "WARNING", "Line duplicates an earlier term (after normalization).",
     "Relativity ignores duplicates on load; removed from the clean list.", "—",
     "Adopt as-is?"),
    ("IMPLICIT_PHRASE", "INFO", "Adjacent words with no operator.",
     "dtSearch runs them as an exact phrase, not an AND. Quoted for clarity; use AND if both-anywhere was intended.",
     'market analysis  →  "market analysis"',
     "Do requestors usually MEAN phrase or AND? Add a line to the intake template asking."),
    ("OPERATOR_CASE", "INFO", "Operators in lowercase.",
     "Recognized either way; capitalized by convention so operators are visually distinct from literal words.",
     "a and b  →  a AND b",
     "Adopt as-is?"),
    ("NOT_WN_ASYMMETRY", "INFO", "Query uses NOT w/N.",
     "NOT w/N is not symmetrical — order matters. Confirm direction.",
     "a NOT w/5 b  ≠  b NOT w/5 a",
     "Adopt as-is (reminder only)?"),
    ("POSSIBLE_MISSPELLING", "INFO", "Words not found in the dictionary (optional; needs pyspellchecker).",
     "No auto-change — misspellings are sometimes deliberate in eDiscovery (they catch the same typo in documents). Suggestions listed for the requestor.",
     "contrct  →  suggestion: contract? (no change made)",
     "Keep this check on by default? Who reviews the suggestions with the case team?"),
    ("NOISE_WORD_IN_PHRASE / WHITESPACE / CONTROL_CHARS / EMPTY_LINE", "INFO", "Hygiene notes.",
     "Phrase noise words act as one-word placeholders; whitespace/control characters trimmed; blank lines skipped.", "—",
     "Adopt as-is?"),
]

TOPICS = [
    ("Index configuration", "Get written confirmation of the noise-word list and alphabet file for each active workspace; update the script constants where customized. Several rules are only advisory until this is done."),
    ("Thresholds", "Agree the tunable thresholds as team standards: PROX_TOO_WIDE (25 words), MANY_PROX_OPS (3 operators), OVERBROAD_TERM (3 characters), DEFAULT_PROX_DISTANCE (5)."),
    ("Comma-list house rule", "Decide the default fix for comma-separated lines: split to separate lines (per-term hit counts) vs join with OR (single count)."),
    ("Approval boundary", "Confirm the split between syntax-only fixes (notify) and meaning-changing revisions (written approval) matches how the case teams want to work."),
    ("Assumed defaults", "The script applies labeled assumptions so terms can run: w/5 for a missing distance, w/1 for w/0, removing a mid-word #. Ratify these as house defaults or change them."),
    ("Metadata house rule", "field:value conditions are extracted to the 'Extracted Conditions' sheet and belong in the searchable-set saved search. Agree who configures them and whether subject: phrases are re-added as content terms."),
    ("Self-verification gate", "Agree that only CLEAN and ADVISORY terms (Revised Term Check column) may be loaded, and NEEDS REVIEW always goes back to a human."),
    ("Script ownership", "Name an owner for str_term_validator.py: change requests, version control, and communicating rule changes to the team."),
    ("Artifact storage", "Agree the folder standard and where the request log lives."),
    ("Training", "Decide how new analysts learn this: walkthrough with examples/sample_terms.txt, then one supervised live matter?"),
]


def build_xlsx(path):
    wb = Workbook()

    # ---------------- How to use ----------------
    ws = wb.active
    ws.title = "How to use"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 100
    ws.column_dimensions["C"].width = 12

    def put(row, text, bold=False, size=10, color="222222"):
        c = ws.cell(row=row, column=2, value=text)
        c.font = Font(name="Arial", size=size, bold=bold, color=color)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    put(2, "STR search-term validation rules — team review", bold=True, size=14, color=HEAD_BG)
    put(4, "Purpose: walk through the validator's rule set together, agree how we handle each case as a team, "
           "and capture owners for the follow-ups. One row per rule on the 'Rule review' sheet; broader topics "
           "that cut across rules are on 'Open topics'.")
    put(6, "How to work the 'Rule review' sheet in the meeting:", bold=True)
    put(7, "•  Columns A–F describe the rule and are read-only: the use case, severity, what it detects, how the "
           "script currently handles it, an example, and a suggested discussion point.")
    put(8, "•  Fill in the YELLOW columns (G–J) as you go: Team decision (dropdown), Owner, Notes, and Status. "
           "Row 3 is a filled-in example — grey, not a real rule.")
    put(10, "Team decision options:", bold=True)
    put(11, "•  Adopt as-is — the script's current handling becomes our standard.\n"
            "•  Adopt with changes — agree the change in Notes; owner files a script change request.\n"
            "•  Needs follow-up — can't decide today; owner takes the action in Notes.\n"
            "•  Reject — we won't enforce this rule; note why.")
    put(13, "Status counts (update as you fill in the sheet):", bold=True)
    for i, (label, formula) in enumerate([
        ("Rules total", "=COUNTA('Rule review'!A4:A200)"),
        ("Agreed", "=COUNTIF('Rule review'!J4:J200,\"Agreed\")"),
        ("Open", "=COUNTIF('Rule review'!J4:J200,\"Open\")"),
        ("Parked", "=COUNTIF('Rule review'!J4:J200,\"Parked\")"),
    ]):
        ws.cell(row=14 + i, column=2, value=label).font = F()
        ws.cell(row=14 + i, column=3, value=formula).font = F(bold=True)
    put(19, "Severity colors: red = ERROR (must be fixed before loading an STR), amber = WARNING (needs a recorded "
            "decision), blue = INFO (housekeeping). The validator auto-repairs terms to a stable state and "
            "re-validates its own output (report column 'Revised Term Check': CLEAN / ADVISORY / NEEDS REVIEW). "
            "Source: str_term_validator.py and the Relativity dtSearch / Search Terms Report documentation.")

    # ---------------- Rule review ----------------
    ws = wb.create_sheet("Rule review")
    headers = ["Use cases", "Sev.", "What it detects", "Handling (current script behavior)",
               "Example (original → revised)", "Discussion point",
               "Team decision", "Owner", "Notes / agreed change", "Status"]
    widths = [26, 9, 30, 40, 30, 36, 16, 12, 30, 9]
    ws.cell(row=1, column=1, value="Rule-by-rule review — fill the yellow columns in the meeting").font = \
        F(bold=True, size=12, color=HEAD_BG)
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=2, column=i, value=h)
        c.font, c.fill, c.border = HEADF, PatternFill("solid", fgColor=HEAD_BG), BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w

    ex = ["(EXAMPLE) UNBALANCED_QUOTES", "ERROR", "Odd number of double quotes.",
          "Auto-balances the quotes.", '"unclosed  →  "unclosed"',
          "Trust auto-balance or confirm with requestor?",
          "Adopt with changes", "A. Analyst", "Auto-fix OK, but memo must always show before/after.", "Agreed"]
    for i, val in enumerate(ex, start=1):
        c = ws.cell(row=3, column=i, value=val)
        c.font = F(italic=True, color="666666")
        c.fill = PatternFill("solid", fgColor=EXAMPLE_BG)
        c.alignment, c.border = Alignment(wrap_text=True, vertical="top"), BORDER

    row = 4
    for code, sev, det, handle, exm, disc in RULES:
        for i, val in enumerate([code, sev, det, handle, exm, disc, "", "", "", "Open"], start=1):
            c = ws.cell(row=row, column=i, value=val)
            c.font = F()
            c.alignment, c.border = Alignment(wrap_text=True, vertical="top"), BORDER
            if i == 2:
                c.fill = PatternFill("solid", fgColor=SEV_FILL[sev])
                c.font = F(bold=True)
            elif i >= 7:
                c.fill = PatternFill("solid", fgColor=FILLIN)
        row += 1
    last = row - 1
    dv_decision = DataValidation(type="list", allow_blank=True,
                                 formula1='"Adopt as-is,Adopt with changes,Needs follow-up,Reject"')
    dv_status = DataValidation(type="list", formula1='"Open,Agreed,Parked"', allow_blank=True)
    ws.add_data_validation(dv_decision)
    ws.add_data_validation(dv_status)
    dv_decision.add(f"G4:G{last}")
    dv_status.add(f"J4:J{last}")
    ws.freeze_panes = "C4"
    ws.auto_filter.ref = f"A2:J{last}"

    # ---------------- Open topics ----------------
    ws = wb.create_sheet("Open topics")
    ws.cell(row=1, column=1, value="Cross-cutting topics — not tied to a single rule").font = \
        F(bold=True, size=12, color=HEAD_BG)
    for i, (h, w) in enumerate(zip(["Topic", "What we need to agree", "Owner", "Due", "Status", "Notes"],
                                   [20, 60, 12, 12, 10, 30]), start=1):
        c = ws.cell(row=2, column=i, value=h)
        c.font, c.fill, c.border = HEADF, PatternFill("solid", fgColor=HEAD_BG), BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    r = 3
    for topic, ask in TOPICS:
        for i, val in enumerate([topic, ask, "", "", "Open", ""], start=1):
            c = ws.cell(row=r, column=i, value=val)
            c.font = F()
            c.alignment, c.border = Alignment(wrap_text=True, vertical="top"), BORDER
            if i >= 3:
                c.fill = PatternFill("solid", fgColor=FILLIN)
        r += 1
    dv2 = DataValidation(type="list", formula1='"Open,Agreed,Parked"', allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add(f"E3:E{r - 1}")
    ws.freeze_panes = "A3"

    wb.save(path)


def build_markdown(path):
    def esc(s):
        return s.replace("|", "\\|")
    lines = [
        "# Validation rule catalog",
        "",
        "Every finding `str_term_validator.py` can raise. **ERROR** must be resolved before a term is "
        "loaded into an STR, **WARNING** needs a recorded decision, **INFO** is a note. Most findings are "
        "auto-repaired; the report's *Revised Term Check* column shows whether the repaired term "
        "re-validates CLEAN, ADVISORY, or NEEDS REVIEW.",
        "",
        "| Code | Severity | What it detects | Handling | Example (original → revised) |",
        "|---|---|---|---|---|",
    ]
    for code, sev, det, handle, exm, _ in RULES:
        lines.append(f"| `{esc(code)}` | {sev} | {esc(det)} | {esc(handle)} | {esc(exm)} |")
    lines += ["", "_Generated by `docs/build/build_review_tracker.py`._", ""]
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))


if __name__ == "__main__":
    build_xlsx(os.path.join(DOCS, "STR_Validation_Rules_Team_Review.xlsx"))
    build_markdown(os.path.join(DOCS, "RULES.md"))
    print(f"rules: {len(RULES)}  topics: {len(TOPICS)}")
