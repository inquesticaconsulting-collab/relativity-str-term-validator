#!/usr/bin/env python3
"""
str_term_validator.py - Relativity Search Terms Report (STR) term validator
===========================================================================

Validates proposed search terms before they are loaded into a Relativity
Search Terms Report. Per Relativity documentation, every line entered in an
STR runs as an individual dtSearch query, so each term is checked against
dtSearch syntax rules plus the STR-specific constraints.

What it does, per term:
  1. Normalizes pasted text (smart quotes, dashes, hidden characters).
  2. Extracts metadata conditions written as field:value (custodian:, date:,
     ext:, ...). These are not dtSearch syntax - they belong in the saved
     search that defines the STR's searchable set - and their values are
     validated (impossible dates, misspelled extensions and field names).
  3. Runs 40+ checks and auto-repairs every fixable defect.
  4. Re-validates the revised term until it stops changing (fixpoint), then
     self-verifies it: CLEAN / ADVISORY / NEEDS REVIEW.
  5. Labels every default it had to choose as "ASSUMPTION APPLIED".

Input : .xlsx / .csv / .txt - one term per line/row. In spreadsheets a column
        headed "Revised Term", "Term", "Search Term" or "Keyword" is detected
        automatically on any sheet (so a report produced by this script can be
        fed straight back in); otherwise the first column is used.
Output: color-coded Excel report (Summary, Validation Detail, Extracted
        Conditions) and, with --clean, an STR-ready .txt term list.

Usage:
    python str_term_validator.py terms.xlsx
    python str_term_validator.py terms.csv -o report.xlsx --clean clean.txt

Requires: openpyxl.  Optional: pyspellchecker (spelling suggestions).

Rule sources (Relativity documentation, RelativityOne help):
  - "Using dtSearch syntax options"
  - "Create and edit search term reports"
  - "Search terms reports"

NOTE: noise-word and alphabet-file behavior is index-specific. This script
assumes Relativity's DEFAULT dtSearch index settings; adjust NOISE_WORDS and
the special-character guidance below if your index is customized.

Not affiliated with or endorsed by Relativity ODA LLC. Not legal advice.

Copyright (c) 2026 Inquestica Consulting - Inhaber: Neeraj Khandelwal
SPDX-License-Identifier: MIT
"""

import argparse
import calendar
import csv
import os
import re
import unicodedata
from dataclasses import dataclass, field

__version__ = "2.0.0"

# ---------------------------------------------------------------------------
# Constants derived from Relativity / dtSearch documentation
# ---------------------------------------------------------------------------

# STR: "a single term has a character limit of 450"
STR_MAX_TERM_LENGTH = 450

# Default dtSearch noise-word list (indexed but NOT searchable unless the
# list is modified and the index rebuilt). Admins can customize per index.
NOISE_WORDS = {
    "a", "about", "after", "all", "also", "an", "and", "another", "any",
    "are", "as", "at", "be", "because", "been", "before", "being", "between",
    "both", "but", "by", "came", "can", "come", "could", "did", "do", "does",
    "each", "else", "for", "from", "get", "got", "had", "has", "have", "he",
    "her", "here", "him", "himself", "his", "how", "if", "in", "into", "is",
    "it", "its", "just", "like", "make", "many", "me", "might", "more",
    "most", "much", "must", "my", "never", "no", "now", "of", "on", "only",
    "or", "other", "our", "out", "over", "re", "said", "same", "see",
    "should", "since", "so", "some", "still", "such", "take", "than", "that",
    "the", "their", "them", "then", "there", "these", "they", "this",
    "those", "through", "to", "too", "under", "up", "use", "very", "want",
    "was", "way", "we", "well", "were", "what", "when", "where", "which",
    "while", "who", "will", "with", "would", "you", "your",
}

# dtSearch reserved connector words - quote the phrase to search them literally
RESERVED_WORDS = {"and", "or", "not", "to", "contains"}

# Built-in positional tokens present in every index
POSITIONAL_TOKENS = {"xfirstword", "xlastword"}

# Proximity operator token, e.g. w/5, pre/10  (NOT w/5 = NOT token + w/5)
PROX_VALID_RE = re.compile(r"^(w|pre)/(\d+)$", re.IGNORECASE)
# Anything that *looks like* a proximity operator (to catch malformed ones)
PROX_LOOKALIKE_RE = re.compile(r"^(w|pre)/\S*$", re.IGNORECASE)

# Tokenizer: quoted phrase | parenthesis | run of non-space
TOKEN_RE = re.compile(r'"[^"]*"|\(|\)|[^\s()]+')

# Characters copied from Word / Outlook that break dtSearch queries
UNICODE_FIXES = {
    "\u2018": "'", "\u2019": "'",       # curly single quotes
    "\u201C": '"', "\u201D": '"',       # curly double quotes
    "\u2013": "-", "\u2014": "-",       # en / em dash
    "\u00A0": " ", "\u2007": " ",       # non-breaking spaces
    "\u202F": " ",
    "\u200B": "", "\uFEFF": "",         # zero-width space / BOM
    "\u2026": "...",                    # ellipsis
}

# Thresholds for best-practice warnings (team-tunable)
PROX_COUNT_WARN = 3       # "large numbers of proximity operators" per query
PROX_DISTANCE_WARN = 25   # very wide proximity window
SHORT_TERM_LEN = 3        # standalone terms this short are usually overbroad
DEFAULT_PROX_DISTANCE = 5  # assumed when "w/" has no number (flagged)

# Metadata field names commonly (and wrongly) written as field:value inside
# STR terms. These belong in the searchable-set saved search, not the term.
KNOWN_FIELDS = ["custodian", "author", "subject", "date", "ext", "fileext",
                "filetype", "filename", "folder", "project", "title", "from",
                "to", "cc", "bcc", "doctype", "recipient", "sender"]
KNOWN_EXTS = ["pdf", "doc", "docx", "xls", "xlsx", "xlsm", "csv", "msg",
              "eml", "ppt", "pptx", "txt", "zip", "htm", "html", "rtf"]

# field:value or field:"quoted value", optionally "TO <value>" (date ranges)
FIELD_RE = re.compile(
    r'\b([A-Za-z][A-Za-z_]{1,20}):\s*("[^"]*"|[^\s()]+)'
    r'(\s+(?i:to)\s+("[^"]*"|[^\s()]+))?')

# Empty function call dtSearch doesn't have, e.g. stemming()
FUNC_RE = re.compile(
    r'\b(?!date\b|mail\b|creditcard\b)([A-Za-z_]\w*)\(\s*\)')

SEV_ERROR, SEV_WARNING, SEV_INFO = "ERROR", "WARNING", "INFO"
SEV_RANK = {SEV_ERROR: 3, SEV_WARNING: 2, SEV_INFO: 1}

SPELLCHECK_ENABLED = True


@dataclass
class Issue:
    severity: str
    code: str
    message: str      # what is wrong
    reasoning: str    # why it matters / what the revision does


@dataclass
class Result:
    line_no: int
    original: str
    revised: str = ""
    issues: list = field(default_factory=list)
    conditions: list = field(default_factory=list)  # extracted metadata
    residual: str = ""  # self-verification result for the revised term

    @property
    def status(self):
        if not self.issues:
            return "VALID"
        worst = max(self.issues, key=lambda i: SEV_RANK[i.severity])
        return worst.severity

    def add(self, severity, code, message, reasoning):
        self.issues.append(Issue(severity, code, message, reasoning))


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------

def _osa_distance(a, b):
    """Optimal-string-alignment edit distance (counts a transposition as 1)."""
    la, lb = len(a), len(b)
    d = [[0] * (lb + 1) for _ in range(la + 1)]
    for i in range(la + 1):
        d[i][0] = i
    for j in range(lb + 1):
        d[0][j] = j
    for i in range(1, la + 1):
        for j in range(1, lb + 1):
            cost = 0 if a[i - 1] == b[j - 1] else 1
            d[i][j] = min(d[i - 1][j] + 1, d[i][j - 1] + 1,
                          d[i - 1][j - 1] + cost)
            if (i > 1 and j > 1 and a[i - 1] == b[j - 2]
                    and a[i - 2] == b[j - 1]):
                d[i][j] = min(d[i][j], d[i - 2][j - 2] + 1)
    return d[la][lb]


def _closest(word, candidates, max_dist=2):
    """Closest candidate within max_dist (ties prefer equal length)."""
    best = None
    for c in candidates:
        dist = _osa_distance(word, c)
        if dist <= max_dist:
            key = (dist, abs(len(c) - len(word)))
            if best is None or key < best[0]:
                best = (key, c)
    return best[1] if best else None


def _split_outside_quotes(term, seps=",;"):
    parts, buf, in_q = [], [], False
    for ch in term:
        if ch == '"':
            in_q = not in_q
        if ch in seps and not in_q:
            parts.append("".join(buf))
            buf = []
        else:
            buf.append(ch)
    parts.append("".join(buf))
    return parts


def _collapse(text):
    return re.sub(r"\s+", " ", text).strip()


# ---------------------------------------------------------------------------
# Input readers
# ---------------------------------------------------------------------------

# Header priority: a revised column beats a generic term column, which beats
# the original column - so a validation report can be fed straight back in.
HEADER_PRIORITY = [
    {"revised term", "revised terms", "final term", "final terms",
     "clean term", "clean terms"},
    {"term", "terms", "search term", "search terms", "searchterm",
     "keyword", "keywords", "str term", "str terms"},
    {"original term", "original terms"},
]


def read_terms(path):
    """Return list of (line_no, raw_term) from txt/csv/xlsx input."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".txt", ".text", ""):
        return _read_txt(path)
    if ext == ".csv":
        return _read_csv(path)
    if ext in (".xlsx", ".xlsm", ".xltx"):
        return _read_xlsx(path)
    raise SystemExit(f"Unsupported input type '{ext}'. Use .txt, .csv or .xlsx")


def _read_txt(path):
    out = []
    with open(path, "r", encoding="utf-8-sig", errors="replace") as fh:
        for i, line in enumerate(fh, start=1):
            out.append((i, line.rstrip("\r\n")))
    return out


def _pick_column(header_row):
    """Index of the best terms column, or None if the row is not a header."""
    cells = [str(c).strip().lower() if c is not None else "" for c in header_row]
    for group in HEADER_PRIORITY:
        for idx, cell in enumerate(cells):
            if cell in group:
                return idx
    return None


def _read_csv(path):
    with open(path, "r", encoding="utf-8-sig", errors="replace",
              newline="") as fh:
        rows = list(csv.reader(fh))
    return _rows_to_terms(rows)


def _read_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit("openpyxl is required: pip install openpyxl")
    wb = load_workbook(path, read_only=True, data_only=True)
    chosen = None
    for ws in wb.worksheets:           # prefer a sheet with a terms header
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if rows and _pick_column(rows[0]) is not None:
            chosen = rows
            break
    if chosen is None:
        chosen = [list(r) for r in wb.active.iter_rows(values_only=True)]
    wb.close()
    return _rows_to_terms(chosen)


def _rows_to_terms(rows):
    if not rows:
        return []
    col = _pick_column(rows[0])
    start = 0
    if col is None:
        col = 0
    else:
        start = 1  # skip the header row
    out = []
    for i, row in enumerate(rows[start:], start=start + 1):
        val = row[col] if col < len(row) else None
        out.append((i, "" if val is None else str(val)))
    return out


# ---------------------------------------------------------------------------
# Normalization (auto-fixes applied to the revised term)
# ---------------------------------------------------------------------------

def normalize(term, res):
    fixed = term
    replaced = set()
    for bad, good in UNICODE_FIXES.items():
        if bad in fixed:
            replaced.add(f"U+{ord(bad):04X}")
            fixed = fixed.replace(bad, good)
    if replaced:
        res.add(SEV_WARNING, "UNICODE_CHARS",
                f"Non-standard characters found ({', '.join(sorted(replaced))}) - "
                "typically smart quotes/dashes pasted from Word or Outlook.",
                "dtSearch does not treat curly quotes as phrase delimiters and "
                "special dashes/spaces are not indexed as typed. Replaced with "
                "plain ASCII equivalents so the query behaves as intended.")
    cleaned = "".join(" " if unicodedata.category(ch).startswith("C") else ch
                      for ch in fixed)
    if cleaned != fixed:
        res.add(SEV_INFO, "CONTROL_CHARS",
                "Hidden control characters (tabs, etc.) removed.",
                "Control characters are treated as spaces or ignored by the "
                "index and add no search value.")
        fixed = cleaned
    collapsed = _collapse(fixed)
    if collapsed != fixed:
        res.add(SEV_INFO, "WHITESPACE",
                "Leading/trailing/multiple spaces trimmed.",
                "Extra whitespace is not meaningful in dtSearch and makes "
                "duplicate detection unreliable (STR ignores duplicates by "
                "exact text).")
    return collapsed


# ---------------------------------------------------------------------------
# Tokenization helpers
# ---------------------------------------------------------------------------

def tokenize(term):
    """Split a term into tokens: quoted phrases, parens, words/operators."""
    return TOKEN_RE.findall(term)


def classify(token):
    """Classify a token: OPEN, CLOSE, PHRASE, NOT, BINOP, PROX, WORD."""
    if token == "(":
        return "OPEN"
    if token == ")":
        return "CLOSE"
    if token.startswith('"'):
        return "PHRASE"
    low = token.lower()
    if low == "not":
        return "NOT"
    if low in ("and", "or", "to", "contains"):
        return "BINOP"
    if PROX_VALID_RE.match(token):
        return "PROX"
    return "WORD"


def words_in_phrase(phrase_token):
    """Words inside a quoted phrase token."""
    inner = phrase_token.strip('"')
    return re.findall(r"[A-Za-z0-9*?=~%#']+", inner)


def _join_tokens(tokens):
    out = ""
    for t in tokens:
        if t == ")":
            out = out.rstrip() + ") "
        else:
            out += t + " "
    out = _collapse(out)
    return out.replace("( ", "(").replace(" )", ")")


# ---------------------------------------------------------------------------
# Checks - metadata, functions, decoration (run first)
# ---------------------------------------------------------------------------

def _check_date_value(val):
    """Return a note if val looks like a date but is impossible; else None."""
    v = val.strip('"')
    m = re.match(r"^(\d{4})[-/](\d{1,2})[-/](\d{1,2})$", v)
    if not m:
        return None
    y, a, b = int(m.group(1)), int(m.group(2)), int(m.group(3))
    for mo, dy in ((a, b), (b, a)):    # accept YYYY-MM-DD or YYYY-DD-MM
        if 1 <= mo <= 12 and 1 <= dy <= calendar.monthrange(y, mo)[1]:
            return None
    return f"'{v}' is not a possible calendar date - confirm the intended date"


def check_field_syntax(term, res):
    """field:value metadata conditions are not dtSearch content syntax."""
    found = []

    def _repl(m):
        if term[:m.start()].count('"') % 2 == 1:   # inside a quoted phrase
            return m.group(0)
        fieldname, value, torange = m.group(1), m.group(2), m.group(3) or ""
        if value.startswith("//"):                  # a URL like http://...
            return m.group(0)
        notes = []
        fl = fieldname.lower()
        if fl not in KNOWN_FIELDS:
            close = _closest(fl, KNOWN_FIELDS)
            notes.append(f"field name '{fieldname}' - did you mean '{close}'?"
                         if close else
                         f"field name '{fieldname}' not recognized")
        if fl in ("ext", "fileext", "filetype"):
            v = value.strip('".*').lower()
            if v not in KNOWN_EXTS:
                close = _closest(v, KNOWN_EXTS)
                notes.append(f"extension '{v}' looks misspelled - did you "
                             f"mean '{close}'?" if close else
                             f"extension '{v}' is not a common file type - "
                             "verify")
        if fl == "date":
            for v in [value] + ([torange.split()[-1]] if torange else []):
                note = _check_date_value(v)
                if note:
                    notes.append(note)
        found.append({"text": m.group(0), "field": fieldname,
                      "note": "; ".join(notes)})
        return " "

    fixed = FIELD_RE.sub(_repl, term)
    if found:
        res.conditions.extend(found)
        conds = ", ".join(f["text"] for f in found[:6])
        more = "" if len(found) <= 6 else f" (and {len(found) - 6} more)"
        res.add(SEV_ERROR, "FIELD_SYNTAX_UNSUPPORTED",
                f"Metadata conditions written as field:value ({conds}{more}).",
                "field:value is not dtSearch syntax - in a Relativity STR the "
                "whole line searches document CONTENT only. Metadata "
                "restrictions (custodian, date, file extension, author, "
                "folder, project, subject) belong in the saved search that "
                "defines the STR's searchable set, or in index conditions. "
                "They were removed from the term; apply them to the saved "
                "search instead - see the Extracted Conditions sheet. If a "
                "subject phrase should also be searched as document text, add "
                "it back as a quoted phrase.")
        return _collapse(fixed)
    return term


def check_functions(term, res):
    """Unknown empty function calls like stemming() are not dtSearch."""
    hits = FUNC_RE.findall(term)
    if hits:
        res.add(SEV_ERROR, "UNSUPPORTED_FUNCTION",
                "Function-style syntax not supported by dtSearch: "
                f"{', '.join(h + '()' for h in hits[:4])}.",
                "dtSearch has no such function (only date(), mail() and "
                "creditcard() exist). Left in place it would search the bare "
                "word literally. Removed; for stemming, append ~ to the root "
                "word (e.g. apply~).")
        return _collapse(FUNC_RE.sub(" ", term))
    return term


def check_stray_tildes(term, res):
    """Tildes that are neither stemming (word~) nor a numeric range (N~~M)."""
    out, hit, i = [], False, 0
    while i < len(term):
        if term[i] != "~":
            out.append(term[i])
            i += 1
            continue
        j = i
        while j < len(term) and term[j] == "~":
            j += 1
        run = j - i
        before = term[i - 1] if i > 0 else ""
        after = term[j] if j < len(term) else ""
        if run == 2 and (before.isdigit() or after.isdigit()):
            out.append(term[i:j])      # numeric range (validated later)
        elif run == 1 and (before.isalnum() or after.isalnum()):
            out.append("~")            # stemming suffix, or leading ~ that
            #                            check_modifier_placement relocates
        else:
            hit = True                 # stray decoration - drop it
        i = j
    if hit:
        res.add(SEV_ERROR, "STRAY_TILDE",
                "Misplaced tilde(s) (e.g. ~~ or ~~~ attached to a word or "
                "quoted phrase).",
                "~~ is only valid as a numeric range between integers "
                "(12~~24), and a single ~ only as a stemming suffix on a bare "
                "word (apply~) - stemming does not apply to quoted phrases. "
                "The stray tildes were removed; if stemming was intended, "
                "unquote the word and append a single ~.")
        return _collapse("".join(out))
    return term


def check_prox_repair(term, res):
    """Repairable proximity defects: w/7/ and a missing distance."""
    fixed = term
    pat = re.compile(r"\b(w|pre)/(\d+)/+", re.IGNORECASE)
    if pat.search(fixed):
        fixed = pat.sub(lambda m: f"{m.group(1).lower()}/{m.group(2)}", fixed)
        res.add(SEV_ERROR, "PROX_TRAILING_SLASH",
                "Proximity operator with a trailing slash (e.g. w/7/).",
                "dtSearch expects exactly w/N. The extra slash makes the "
                "operator unrecognizable, so the terms would be searched as "
                "an implicit phrase instead. Corrected to w/N.")
    # bare "w/" or "pre/" with no number anywhere after it (spacing such as
    # "w/ 5" is left for check_proximity_form)
    pat2 = re.compile(r"\b(w|pre)/(?!\s*\d)", re.IGNORECASE)
    if pat2.search(fixed):
        fixed = _collapse(pat2.sub(
            lambda m: f"{m.group(1).lower()}/{DEFAULT_PROX_DISTANCE} ", fixed))
        res.add(SEV_WARNING, "PROX_DISTANCE_ASSUMED",
                "Proximity operator had no distance (w/ with no number).",
                f"ASSUMPTION APPLIED: a default distance of "
                f"{DEFAULT_PROX_DISTANCE} words was inserted so the term can "
                "run. Confirm the intended distance with the case team - it "
                "was chosen as a common default, not taken from the source.")
    return fixed


def check_star_runs(term, res):
    """Collapse ** runs - they multiply matches without adding meaning."""
    if re.search(r"\*{2,}", term):
        res.add(SEV_WARNING, "MULTI_WILDCARD_RUN",
                "Consecutive wildcards (**) found.",
                "A run of asterisks behaves like a single * but signals a "
                "typo and can confuse review. Collapsed to one *.")
        return re.sub(r"\*{2,}", "*", term)
    return term


# ---------------------------------------------------------------------------
# Checks - structure and STR rules
# ---------------------------------------------------------------------------

def check_length(term, res):
    if len(term) > STR_MAX_TERM_LENGTH:
        res.add(SEV_ERROR, "OVER_450_CHARS",
                f"Term is {len(term)} characters; the STR limit is "
                f"{STR_MAX_TERM_LENGTH} per term.",
                "Relativity rejects STR terms over 450 characters. Split this "
                "into several shorter terms (each line in an STR is its own "
                "dtSearch query, and multiple narrow queries are also a "
                "documented performance best practice).")
    return term


def check_balanced(term, res):
    fixed = term
    if fixed.count('"') % 2 != 0:
        if fixed.count('"') == 1:
            fixed = fixed.replace('"', "")
            action = "the stray quote was removed"
        else:
            fixed = fixed + '"'
            action = "a closing quote was appended"
        res.add(SEV_ERROR, "UNBALANCED_QUOTES",
                "Odd number of double quotes - the phrase is never closed.",
                "dtSearch treats quotes as phrase delimiters; an unclosed "
                f"quote makes the query invalid or unpredictable ({action} - "
                "review that the intended phrase is quoted correctly).")
    depth, extra_close = 0, 0
    for ch in fixed:
        if ch == "(":
            depth += 1
        elif ch == ")":
            if depth == 0:
                extra_close += 1
            else:
                depth -= 1
    if depth or extra_close:
        if extra_close:
            out, d = [], 0
            for ch in fixed:
                if ch == "(":
                    d += 1
                elif ch == ")":
                    if d == 0:
                        continue
                    d -= 1
                out.append(ch)
            fixed = "".join(out)
        if depth:
            fixed = fixed + ")" * depth
        res.add(SEV_ERROR, "UNBALANCED_PARENS",
                "Parentheses are not balanced.",
                "Unbalanced parentheses make the grouping ambiguous or the "
                "query invalid. Missing closers were appended / stray closers "
                "removed - verify the grouping still matches your intent.")
    if "()" in re.sub(r"\s+", "", fixed):
        fixed = _collapse(re.sub(r"\(\s*\)", " ", fixed))
        res.add(SEV_ERROR, "EMPTY_PARENS",
                "Empty parentheses '()' found.",
                "An empty group contains nothing to search and invalidates "
                "the query; it was removed.")
    return fixed


def check_separators(term, res):
    """Commas/semicolons outside quotes usually mean several terms pasted
    onto a single row."""
    masked = re.sub(r'"[^"]*"', '""', term)
    if not ("," in masked or ";" in masked) or re.search(r"\d,\d", masked):
        return term
    pieces = [p.strip() for p in _split_outside_quotes(term) if p.strip()]
    if len(pieces) > 1:
        fixed = " OR ".join(
            f'"{p}"' if " " in p and not p.startswith('"') and
            not re.search(r"\b(and|or|not|w/\d+|pre/\d+)\b", p, re.I)
            else p
            for p in pieces)
        res.add(SEV_ERROR, "COMMA_SEPARATED_LIST",
                "Line contains comma/semicolon-separated values - dtSearch "
                "treats commas as spaces, so this would run as one exact "
                "phrase, not a list of terms.",
                "In an STR every line is ONE dtSearch query. 'apple, pear' "
                "searches the phrase 'apple pear' and would hit almost "
                "nothing. Rewritten with OR; better still, put each term "
                "on its own line so the report counts hits per term.")
        return fixed
    if pieces:
        res.add(SEV_INFO, "STRAY_SEPARATOR",
                "Leading/trailing comma or semicolon removed.",
                "A dangling separator carries no meaning in dtSearch.")
        return pieces[0]
    return term


def check_operator_grammar(term, res):
    """Sequence validation and repair: dangling and doubled operators."""
    tokens = tokenize(term)
    kinds = [classify(t) for t in tokens]
    if not tokens:
        return term
    if kinds[0] in ("BINOP", "PROX"):
        res.add(SEV_ERROR, "LEADING_OPERATOR",
                f"Query starts with operator '{tokens[0]}'.",
                "A dtSearch query cannot begin with a binary operator - there "
                "is nothing on its left side. Removed; if a term was lost in "
                "copy/paste, restore it.")
    if kinds[-1] in ("BINOP", "PROX", "NOT"):
        res.add(SEV_ERROR, "TRAILING_OPERATOR",
                f"Query ends with operator '{tokens[-1]}'.",
                "The operator has no right-hand term, which invalidates the "
                "query. Removed; if a term was lost in copy/paste, restore "
                "it.")
    prev = None
    for i, k in enumerate(kinds):
        if prev in ("BINOP", "PROX") and k in ("BINOP", "PROX"):
            res.add(SEV_ERROR, "DOUBLE_OPERATOR",
                    f"Consecutive operators '{tokens[i-1]} {tokens[i]}'.",
                    "Two binary operators in a row leave one side empty. A "
                    "single operator is kept (a proximity operator wins over "
                    "AND/OR, since it carries the more specific intent).")
        if prev == "OPEN" and k in ("BINOP", "PROX"):
            res.add(SEV_ERROR, "OPERATOR_AFTER_PAREN",
                    f"Operator '{tokens[i]}' immediately after '('.",
                    "A group cannot start with a binary operator; removed.")
        if prev in ("BINOP", "PROX", "NOT") and k == "CLOSE":
            res.add(SEV_ERROR, "OPERATOR_BEFORE_PAREN",
                    f"Operator '{tokens[i-1]}' immediately before ')'.",
                    "A group cannot end with an operator; removed.")
        if prev == "NOT" and k == "BINOP":
            res.add(SEV_ERROR, "NOT_BEFORE_OPERATOR",
                    f"'NOT {tokens[i]}' - NOT directly followed by an "
                    "operator.",
                    "'NOT OR' / 'NOT AND' is not valid. The NOT was dropped "
                    "and the operator kept - but this pattern often means "
                    "'AND NOT' was intended, which excludes instead of "
                    "including. Confirm with the term's author.")
        prev = k

    # repair
    fixed, fkinds = list(tokens), list(kinds)
    while fixed and fkinds[0] in ("BINOP", "PROX"):
        fixed.pop(0)
        fkinds.pop(0)
    while fixed and fkinds[-1] in ("BINOP", "PROX", "NOT"):
        fixed.pop()
        fkinds.pop()
    out, okinds = [], []
    for t, k in zip(fixed, fkinds):
        if out and okinds[-1] in ("BINOP", "PROX") and k in ("BINOP", "PROX"):
            if okinds[-1] == "BINOP" and k == "PROX":
                out[-1], okinds[-1] = t, k      # 'AND w/10' -> 'w/10'
            continue
        if out and okinds[-1] == "NOT" and k == "BINOP":
            out[-1], okinds[-1] = t, k          # 'NOT OR' -> 'OR'
            continue
        if out and okinds[-1] == "OPEN" and k in ("BINOP", "PROX"):
            continue
        if k == "CLOSE":
            while out and okinds[-1] in ("BINOP", "PROX", "NOT"):
                out.pop()
                okinds.pop()
        out.append(t)
        okinds.append(k)
    if out != tokens:
        return _join_tokens(out)
    return term


def check_proximity_form(term, res):
    """Spacing inside w/N, zero and very wide distances."""
    fixed = term
    spaced = re.compile(r"\b(w|pre)\s+/\s*(\d+)\b|\b(w|pre)/\s+(\d+)\b", re.I)
    if spaced.search(fixed):
        fixed = spaced.sub(lambda m: f"{(m.group(1) or m.group(3)).lower()}/"
                                     f"{m.group(2) or m.group(4)}", fixed)
        res.add(SEV_ERROR, "PROX_SPACING",
                "Space inside a proximity operator (e.g. 'w /5' or 'w/ 5').",
                "dtSearch only recognizes the operator written as w/N or "
                "pre/N with no spaces; otherwise 'w' is searched as a "
                "literal word. Corrected to the joined form.")
    for tok in tokenize(fixed):
        if PROX_LOOKALIKE_RE.match(tok) and not PROX_VALID_RE.match(tok):
            res.add(SEV_ERROR, "PROX_MALFORMED",
                    f"'{tok}' looks like a proximity operator but has no "
                    "valid number.",
                    "Proximity must be w/N or pre/N where N is a positive "
                    "integer (e.g. w/5). Without it, dtSearch will not apply "
                    "proximity. Specify the intended distance.")
        elif PROX_VALID_RE.match(tok):
            n = int(PROX_VALID_RE.match(tok).group(2))
            if n == 0:
                fixed = re.sub(r"\b(w|pre)/0\b",
                               lambda m: m.group(1).lower() + "/1", fixed,
                               flags=re.I)
                res.add(SEV_WARNING, "PROX_ZERO",
                        f"'{tok}' uses a distance of 0.",
                        "ASSUMPTION APPLIED: w/0 is meaningless, so it was "
                        "replaced with w/1 (terms within one word, either "
                        "order). If an exact ordered phrase was intended, "
                        "use a quoted phrase instead - confirm with the "
                        "term's author.")
            elif n > PROX_DISTANCE_WARN:
                res.add(SEV_WARNING, "PROX_TOO_WIDE",
                        f"Very wide proximity window ({tok}).",
                        f"Distances above ~{PROX_DISTANCE_WARN} words often "
                        "behave like an AND across whole paragraphs and "
                        "inflate hit counts. Confirm the width is intended.")
    return fixed


def check_leading_wildcard(term, res):
    """STR: leading wildcards are prohibited."""
    fixed_tokens, hit = [], False
    for tok in tokenize(term):
        if classify(tok) == "WORD" and re.match(r"^[*?=]+\w", tok):
            hit = True
            fixed_tokens.append(tok.lstrip("*?="))
        elif classify(tok) == "WORD" and tok in ("*", "?", "="):
            hit = True
        else:
            fixed_tokens.append(tok)
    if hit:
        res.add(SEV_ERROR, "LEADING_WILDCARD",
                "Leading wildcard (e.g. *term) detected.",
                "Relativity prohibits leading wildcards in Search Terms "
                "Reports, and dtSearch documentation notes wildcards near the "
                "start of a word badly degrade performance. The leading "
                "wildcard was removed - if you need suffix matching, consider "
                "stemming (~) or listing the specific variants with OR.")
        return _join_tokens(fixed_tokens)
    return term


def check_fuzzy(term, res):
    """STR: fuzzy searching (%) is not supported."""
    if "%" in term:
        res.add(SEV_ERROR, "FUZZY_NOT_SUPPORTED",
                "Fuzzy operator '%' found.",
                "Per Relativity documentation, fuzzy searching is not "
                "supported in Search Terms Reports. The '%' was removed. To "
                "capture misspellings, expand the term first with the "
                "dtSearch Dictionary (fuzzy) and paste the resulting "
                "variants joined by OR.")
        return term.replace("%", "")
    return term


def check_modifier_placement(term, res):
    fixed = term
    for tok in tokenize(term):
        if classify(tok) != "WORD":
            continue
        if tok.startswith("~") and not tok.startswith("~~"):
            fixed = fixed.replace(tok, tok.lstrip("~") + "~", 1)
            res.add(SEV_ERROR, "STEM_LEADING",
                    f"Stemming '~' at the start of '{tok}'.",
                    "The stemming operator goes at the END of the root word "
                    "(apply~) - moved there. Note: dtSearch stemming "
                    "supports English only and works from the root form of "
                    "the word.")
        if "#" in tok[1:]:
            fixed = fixed.replace(tok, tok[0] + tok[1:].replace("#", ""), 1)
            res.add(SEV_WARNING, "PHONIC_MISPLACED",
                    f"Phonic '#' inside '{tok}'.",
                    "ASSUMPTION APPLIED: '#' is only valid immediately "
                    "BEFORE a word (#smith) and is not an indexed character "
                    "elsewhere, so it was removed as a likely typo. If a "
                    "phonic search was intended, move '#' to the front of "
                    "the word.")
        if "~~" in tok and not re.match(r"^\d+~~\d+$", tok):
            res.add(SEV_ERROR, "NUMERIC_RANGE_MALFORMED",
                    f"'{tok}' is not a valid numeric range.",
                    "Numeric range syntax is N~~M with integers on both "
                    "sides (e.g. 12~~24). Fix the bounds or remove the "
                    "operator - the intended range cannot be guessed.")
    return fixed


def check_wildcards(term, res):
    tokens = tokenize(term)
    for tok in tokens:
        if classify(tok) != "WORD":
            continue
        if "*" in tok[1:3] and not tok.startswith("*"):
            res.add(SEV_WARNING, "EARLY_WILDCARD",
                    f"Wildcard very close to the start of '{tok}'.",
                    "dtSearch warns that '*' near the beginning of a word "
                    "slows searches dramatically - the engine must scan "
                    "nearly the whole word list. Move the wildcard later or "
                    "spell out variants with OR.")
        if tok.count("*") > 1:
            res.add(SEV_WARNING, "MULTI_WILDCARD",
                    f"Multiple '*' wildcards in '{tok}'.",
                    "Multiple wildcards multiply the candidate words and can "
                    "return far more documents than intended. Tighten to a "
                    "single trailing wildcard or explicit variants.")
        if tok and tok.strip("*?=") == "":
            res.add(SEV_ERROR, "WILDCARD_ONLY",
                    f"'{tok}' is wildcard-only.",
                    "A bare wildcard matches every word in the index and "
                    "will hit essentially all documents; it also risks "
                    "matching the built-in xfirstword/xlastword tokens. "
                    "Replace with a real term.")
        if re.match(r"^x\*", tok, re.I):
            res.add(SEV_WARNING, "X_WILDCARD_TRAP",
                    f"'{tok}' matches the built-in tokens xfirstword/"
                    "xlastword.",
                    "Every document contains xfirstword and xlastword in the "
                    "index, so a wildcard covering them returns ALL "
                    "documents. Narrow the pattern (add more leading "
                    "characters).")
    return term


def check_noise_and_reserved(term, res):
    tokens = tokenize(term)
    kinds = [classify(t) for t in tokens]
    plain_words = [t.lower() for t, k in zip(tokens, kinds) if k == "WORD"]

    if (len(tokens) == 1 and kinds[0] == "WORD"
            and tokens[0].lower().strip("*?=~#") in NOISE_WORDS):
        res.add(SEV_ERROR, "NOISE_WORD_TERM",
                f"'{tokens[0]}' is a dtSearch noise word.",
                "Noise words are not searchable in a default dtSearch index, "
                "so this term will return zero hits. Replace it with a more "
                "specific term, or have an admin remove the word from the "
                "index's noise-word list and rebuild the index.")
        return term

    noisy = sorted({w for w in plain_words
                    if w.strip("*?=~#") in NOISE_WORDS
                    and w not in RESERVED_WORDS})
    if noisy:
        res.add(SEV_WARNING, "NOISE_WORD_IN_QUERY",
                f"Contains dtSearch noise word(s): {', '.join(noisy)}.",
                "Noise words are skipped by the index. In a phrase they act "
                "as a one-word placeholder rather than the literal word "
                "('statement of work' also matches 'statement for work'); "
                "standalone they contribute nothing. Verify the hit counts "
                "make sense, or rephrase without the noise word.")

    for t, k in zip(tokens, kinds):
        if k == "PHRASE":
            ph_noise = sorted({w.lower() for w in words_in_phrase(t)
                               if w.lower() in NOISE_WORDS})
            if ph_noise:
                res.add(SEV_INFO, "NOISE_WORD_IN_PHRASE",
                        f"Quoted phrase {t} contains noise word(s): "
                        f"{', '.join(ph_noise)}.",
                        "Inside a phrase a noise word matches ANY single "
                        "word in that position, which usually still returns "
                        "the intended documents but can slightly over-"
                        "include. Generally acceptable - just be aware.")

    literal = [t for t, k in zip(tokens, kinds)
               if k == "BINOP" and t.lower() in ("to", "contains")]
    if literal:
        res.add(SEV_WARNING, "RESERVED_WORD_UNQUOTED",
                f"Unquoted reserved word(s): {', '.join(literal)}.",
                "'to' and 'contains' are dtSearch connector words. If they "
                "are meant as ordinary words, put the whole phrase in quotes "
                "(e.g. \"go to market\") - unquoted, dtSearch may read them "
                "as operators.")
    return term


def check_redundant_operand(term, res):
    """'deal OR deal' - same word on both sides of an operator."""
    pat = re.compile(r"\b(\S+)\s+(?:or|and)\s+\1\b(?!\S)", re.IGNORECASE)
    fixed, n = term, 0
    while True:
        new = pat.sub(lambda m: m.group(1), fixed)
        if new == fixed:
            break
        fixed, n = new, n + 1
    if n:
        res.add(SEV_WARNING, "REDUNDANT_OPERAND",
                "The same word appears on both sides of AND/OR.",
                "'x OR x' and 'x AND x' are equivalent to just 'x' - the "
                "repetition adds no hits and suggests a copy/paste slip. "
                "Collapsed to the single term; check whether a different "
                "second term was intended.")
        return fixed
    return term


def _make_precedence_explicit(term):
    """Rewrite 'a AND b OR c' as 'a AND (b OR c)' (dtSearch default)."""
    parts = re.split(r"\s+(?i:and)\s+", term)
    if len(parts) < 2:
        return term
    out = []
    for p in parts:
        if re.search(r"\s(?i:or)\s", p) and not p.strip().startswith("("):
            out.append(f"({p.strip()})")
        else:
            out.append(p.strip())
    return " AND ".join(out)


def check_precedence(term, res):
    """Mixed AND/OR without parentheses - dtSearch evaluates OR before AND."""
    tokens = tokenize(term)
    kinds = [classify(t) for t in tokens]
    has_and = any(t.lower() == "and" for t, k in zip(tokens, kinds)
                  if k == "BINOP")
    has_or = any(t.lower() == "or" for t, k in zip(tokens, kinds)
                 if k == "BINOP")
    if has_and and has_or and "(" not in term:
        res.add(SEV_WARNING, "MIXED_AND_OR",
                "AND and OR mixed without parentheses.",
                "dtSearch applies OR before AND, so 'a AND b OR c' runs as "
                "'a AND (b OR c)' - which may not be what you meant. The "
                "revised term adds parentheses showing how dtSearch will "
                "interpret the query as written; confirm the grouping.")
        return _make_precedence_explicit(term)
    return term


def check_proximity_semantics(term, res):
    tokens = tokenize(term)
    prox = [t for t in tokens if classify(t) == "PROX"]
    if not prox:
        return term
    if len(prox) >= PROX_COUNT_WARN:
        res.add(SEV_WARNING, "MANY_PROX_OPS",
                f"{len(prox)} proximity operators in one query.",
                "Relativity's STR guidance: large numbers of proximity "
                "operators in a single query slow report generation. Split "
                "into multiple lines (each is reported separately) if "
                "possible.")
    if len(prox) >= 2:
        res.add(SEV_WARNING, "NESTED_PROX",
                "Multiple/nested proximity operators.",
                "Nesting proximity operators is discouraged in the STR "
                "documentation for performance reasons, and chained "
                "proximity ('a w/5 b w/5 c') can be evaluated ambiguously. "
                "Prefer '(a w/5 b) AND (b w/5 c)' or separate lines.")
    if (re.search(r"\)[^()]*\b(?:w|pre)/\d+", term, re.I)
            or re.search(r"\b(?:w|pre)/\d+[^()]*\(", term, re.I)):
        for grp in re.findall(r"\(([^()]*)\)", term):
            if re.search(r"\s(?i:and)\s", grp):
                res.add(SEV_WARNING, "AND_INSIDE_PROX_GROUP",
                        "A parenthesized group joined by W/N or PRE/N "
                        "contains AND.",
                        "Relativity documents '(a AND b) w/10 (c AND d)' as "
                        "an ambiguous search that should not be used - "
                        "results are unpredictable. Use OR inside groups "
                        "connected by proximity, or restructure as "
                        "'(a w/10 c) AND (b w/10 d)'.")
                break
    if re.search(r"\bnot\s+w/\d+", term, re.I):
        res.add(SEV_INFO, "NOT_WN_ASYMMETRY",
                "Uses NOT W/N.",
                "Unlike W/N, NOT W/N is not symmetrical: 'a NOT w/5 b' "
                "(a's not within 5 words of b) is different from "
                "'b NOT w/5 a'. Double-check the order matches your intent.")
    return term


def check_special_chars(term, res):
    """Characters not in the default searchable alphabet."""
    notes = []
    if "&" in term:
        notes.append(("&", "'&' is not an indexed character by default - "
                      "'AT&T' is indexed as the two words 'at' and 't' (and "
                      "'at' is a noise word). Search a quoted phrase "
                      "(\"at t\"), add '&' to the index alphabet, or use a "
                      "regular expression."))
    if "_" in term:
        notes.append(("_", "The underscore is not recognized by a default "
                      "dtSearch alphabet; behavior depends on the index "
                      "alphabet file. Verify against your index or split "
                      "the words."))
    if ":" in term and not re.search(r"\d:\d", term):
        notes.append((":", "':' is not searchable in content by default - "
                      "Relativity documents using a regular expression to "
                      "find colons in text."))
    stripped = re.sub(r"\b(?:w|pre)/\d+", " ", term, flags=re.I)
    for ch in "@!^|[]{}<>\\/":
        if ch in stripped:
            if ch == "@" and re.search(r"\S+@\S+", term):
                notes.append((ch, "Email addresses: '@' and '.' are word "
                              "separators in a default index, so "
                              "'jdoe@example.com' is indexed as separate "
                              "words. Consider the mail() recognition "
                              "syntax - mail(jdoe@example.com) - or a "
                              "quoted phrase."))
            else:
                notes.append((ch, f"'{ch}' is typically a space/ignored "
                              "character in the default alphabet and will "
                              "not be matched literally."))
    seen = set()
    for ch, msg in notes:
        if ch in seen:
            continue
        seen.add(ch)
        res.add(SEV_WARNING, "SPECIAL_CHAR",
                f"Special character '{ch}' present.",
                msg + " (Alphabet-file behavior is index-specific - confirm "
                "with your Relativity admin.)")
    return term


def check_breadth(term, res):
    tokens = tokenize(term)
    if len(tokens) == 1 and classify(tokens[0]) == "WORD":
        bare = tokens[0].strip("*?=~#")
        if bare and len(bare) <= SHORT_TERM_LEN and not bare.isdigit():
            res.add(SEV_WARNING, "OVERBROAD_TERM",
                    f"Very short standalone term '{tokens[0]}'.",
                    "Short tokens (acronyms aside) match huge volumes and "
                    "false positives - Relativity's guidance is to write "
                    "specific terms and run narrower queries. If this is an "
                    "acronym, consider pairing it with context, e.g. "
                    f"'{bare} w/10 <related-term>'.")
    return term


def check_hyphen(term, res):
    words = [t for t in tokenize(term) if classify(t) == "WORD"]
    hyphenated = [w for w in words
                  if "-" in w.strip("-") and not re.match(r"^\d[\d-]*$", w)]
    if not hyphenated:
        return term
    res.add(SEV_WARNING, "HYPHENATED_TERM",
            f"Hyphenated term(s): {', '.join(hyphenated)}.",
            "Hyphen handling depends on the index's alphabet file (dtSearch "
            "can treat '-' as a space, as searchable, or ignore it - "
            "commonly it behaves like a space, so 'e-mail' matches like the "
            "phrase 'e mail' and does NOT match 'email'). To be safe, "
            "search all variants with OR and confirm the hyphen setting "
            "with your admin.")
    if len(tokenize(term)) == 1:
        w = hyphenated[0]
        return f'("{w}" OR {w.replace("-", "")} OR "{w.replace("-", " ")}")'
    return term


def check_implicit_phrase(term, res):
    """Adjacent unquoted words form an exact phrase - make that explicit."""
    if any(i.code in ("PROX_MALFORMED", "PROX_SPACING") for i in res.issues):
        return term
    tokens = tokenize(term)
    kinds = [classify(t) for t in tokens]
    if len(tokens) > 1 and all(k == "WORD" for k in kinds):
        res.add(SEV_INFO, "IMPLICIT_PHRASE",
                "Multiple words with no operator between them.",
                "dtSearch treats adjacent words as an exact phrase, NOT as "
                "an AND - 'market analysis' only matches the two words "
                "together in order. Quoted for clarity; if you wanted both "
                "words anywhere in the document, use AND instead.")
        return '"' + " ".join(tokens) + '"'
    return term


def check_operator_case(term, res):
    """Uppercase operators outside quotes for readability (cosmetic)."""
    parts = re.split(r'("[^"]*")', term)
    changed = False
    for i, part in enumerate(parts):
        if part.startswith('"'):
            continue
        new = re.sub(r"\b(and|or|not)\b", lambda m: m.group(0).upper(), part)
        new = re.sub(r"\b(w|pre)/(\d+)\b",
                     lambda m: m.group(1).upper() + "/" + m.group(2), new,
                     flags=re.IGNORECASE)
        if new != part:
            changed = True
            parts[i] = new
    if changed:
        res.add(SEV_INFO, "OPERATOR_CASE",
                "Operators written in lowercase.",
                "dtSearch recognizes operators in any case, but convention "
                "is to capitalize them (AND, OR, NOT, W/5) so reviewers can "
                "tell the operator 'and' from the literal word. If you meant "
                "to SEARCH for one of these words, quote the whole phrase.")
        return "".join(parts)
    return term


# ---------------------------------------------------------------------------
# Optional spelling suggestions (never changes the term)
# ---------------------------------------------------------------------------

def check_spelling(res):
    if not SPELLCHECK_ENABLED:
        return
    try:
        from spellchecker import SpellChecker
    except ImportError:
        return
    if not hasattr(check_spelling, "_sp"):
        check_spelling._sp = SpellChecker()
    sp = check_spelling._sp
    words = set()
    for tok in tokenize(res.revised):
        k = classify(tok)
        if k == "PHRASE":
            words.update(words_in_phrase(tok))
        elif k == "WORD":
            words.add(tok)
    cand = {w.lower() for w in words
            if len(w) >= 4 and w.isalpha() and w.lower() not in RESERVED_WORDS}
    unknown = sorted(sp.unknown(cand))[:8]
    if unknown:
        sugg = []
        for w in unknown:
            c = sp.correction(w)
            sugg.append(f"{w} -> {c}?" if c and c != w else w)
        res.add(SEV_INFO, "POSSIBLE_MISSPELLING",
                f"Words not found in the dictionary: {', '.join(unknown)}.",
                "In eDiscovery a misspelled term is sometimes deliberate (it "
                "catches the same typo in documents). If not deliberate, add "
                "the correct spelling with OR so both forms hit "
                f"({'; '.join(sugg)}). No change was made automatically.")


# ---------------------------------------------------------------------------
# Pipeline
# ---------------------------------------------------------------------------

# Order matters: metadata/decoration -> structure -> STR rules -> semantics
# -> style.
CHECK_PIPELINE = [
    check_field_syntax,
    check_functions,
    check_stray_tildes,
    check_prox_repair,
    check_star_runs,
    check_length,
    check_balanced,
    check_separators,
    check_operator_grammar,
    check_proximity_form,
    check_leading_wildcard,
    check_fuzzy,
    check_modifier_placement,
    check_wildcards,
    check_noise_and_reserved,
    check_redundant_operand,
    check_precedence,
    check_proximity_semantics,
    check_special_chars,
    check_breadth,
    check_hyphen,
    check_implicit_phrase,
    check_operator_case,
]

MAX_REPAIR_ITERATIONS = 8


def validate_term(line_no, raw):
    """One pass of every check over a single term."""
    res = Result(line_no=line_no, original=raw)
    term = normalize(raw, res)
    if not term:
        res.add(SEV_INFO, "EMPTY_LINE", "Blank line/cell - skipped.",
                "Empty rows are ignored; nothing to validate.")
        res.revised = ""
        return res
    for check in CHECK_PIPELINE:
        term = check(term, res)
    res.revised = term
    return res


def validate_full(line_no, raw):
    """Validate + repair to a fixpoint, then SELF-VERIFY the revised term.

    The revised term is re-run through the full pipeline until it stops
    changing, so a fix never leaves behind a defect that a second run would
    flag. The final verification pass is reported in res.residual:
      CLEAN            - revised term re-validates with no findings
      ADVISORY: ...    - only warnings/notes remain (the term runs)
      NEEDS REVIEW: ...- an ERROR could not be auto-fixed; human input needed
    """
    res = validate_term(line_no, raw)
    seen = {(i.code, i.message) for i in res.issues}
    current = res.revised
    final_pass = None
    for _ in range(MAX_REPAIR_ITERATIONS):
        nxt = validate_term(line_no, current)
        for c in nxt.conditions:
            if c not in res.conditions:
                res.conditions.append(c)
        for i in nxt.issues:
            key = (i.code, i.message)
            if key not in seen:
                seen.add(key)
                res.issues.append(i)
        if nxt.revised == current:
            final_pass = nxt
            break
        current = nxt.revised
    if final_pass is None:
        final_pass = validate_term(line_no, current)
    res.revised = current

    if not current:
        res.residual = "EMPTY"
        return res
    errs = sorted({i.code for i in final_pass.issues if i.severity == SEV_ERROR})
    warns = sorted({i.code for i in final_pass.issues
                    if i.severity == SEV_WARNING})
    infos = sorted({i.code for i in final_pass.issues if i.severity == SEV_INFO})
    if errs:
        res.residual = "NEEDS REVIEW: " + ", ".join(errs)
    elif warns:
        res.residual = "ADVISORY: " + ", ".join(warns)
    elif infos:
        res.residual = "CLEAN (notes: " + ", ".join(infos) + ")"
    else:
        res.residual = "CLEAN"
    check_spelling(res)
    return res


def flag_duplicates(results):
    seen = {}
    for r in results:
        key = _collapse(r.revised.lower())
        if not key:
            continue
        if key in seen:
            r.add(SEV_WARNING, "DUPLICATE_TERM",
                  f"Duplicate of line {seen[key]} (after normalization).",
                  "Relativity ignores duplicate terms when they are added "
                  "to an STR, so this line adds nothing - remove it to keep "
                  "the term list clean and the report readable.")
        else:
            seen[key] = r.line_no


def validate_all(terms):
    results = [validate_full(no, raw) for no, raw in terms]
    flag_duplicates(results)
    return results


# ---------------------------------------------------------------------------
# Output
# ---------------------------------------------------------------------------

STATUS_ORDER = ["ERROR", "WARNING", "INFO", "VALID"]
STATUS_FILL = {
    "ERROR": "FFF4CCCC", "WARNING": "FFFFF2CC",
    "INFO": "FFD9E7F5", "VALID": "FFD9EAD3",
}


def _xl(value):
    """Strip characters Excel cells cannot hold."""
    if not isinstance(value, str):
        return value
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    return ILLEGAL_CHARACTERS_RE.sub(" ", value)


def write_excel(results, out_path, source_name):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFFFF")
    head_fill = PatternFill("solid", fgColor="FF3D5A80")
    wrap = Alignment(vertical="top", wrap_text=True)

    # ---- Detail sheet ----------------------------------------------------
    ws = wb.active
    ws.title = "Validation Detail"
    ws.append(["#", "Original Term", "Status", "Issue Codes",
               "Findings & Reasoning", "Revised Term",
               "Revised Term Check", "Changed?"])
    for c in ws[1]:
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(vertical="center")
    for r in results:
        codes = ", ".join(dict.fromkeys(i.code for i in r.issues))
        details = "\n".join(
            f"[{i.severity}] {i.message}\n    WHY: {i.reasoning}"
            for i in r.issues) or "No issues found - term is well formed."
        changed = "YES" if r.revised.strip() != r.original.strip() else "no"
        ws.append([_xl(v) for v in (r.line_no, r.original, r.status, codes,
                                    details, r.revised, r.residual, changed)])
        fill = PatternFill("solid", fgColor=STATUS_FILL[r.status])
        for c in ws[ws.max_row]:
            c.fill, c.alignment = fill, wrap
        ws.cell(row=ws.max_row, column=3).font = Font(bold=True)
        chk = ws.cell(row=ws.max_row, column=7)
        chk.font = Font(bold=True)
        for prefix, color in (("CLEAN", "FFD9EAD3"), ("ADVISORY", "FFFFF2CC"),
                              ("NEEDS REVIEW", "FFF4CCCC")):
            if r.residual.startswith(prefix):
                chk.fill = PatternFill("solid", fgColor=color)
    for i, w in enumerate([5, 36, 10, 20, 58, 36, 26, 8], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{ws.max_row}"

    # ---- Extracted Conditions sheet ---------------------------------------
    if any(r.conditions for r in results):
        cs = wb.create_sheet("Extracted Conditions")
        cs.append(["Line #", "Condition (as written)", "Field",
                   "Validation notes", "Where it belongs"])
        for c in cs[1]:
            c.font, c.fill = head_font, head_fill
        for r in results:
            for cond in r.conditions:
                cs.append([_xl(v) for v in (
                    r.line_no, cond["text"], cond["field"],
                    cond.get("note") or "-",
                    "Saved search defining the STR searchable set (or index "
                    "conditions) - not the search term")])
        for col, w in zip("ABCDE", [7, 40, 14, 46, 40]):
            cs.column_dimensions[col].width = w
        for row in cs.iter_rows(min_row=2):
            for c in row:
                c.alignment = wrap
        cs.freeze_panes = "A2"

    # ---- Summary sheet ----------------------------------------------------
    s = wb.create_sheet("Summary", 0)
    s["A1"] = "Relativity STR - Search Term Validation Report"
    s["A1"].font = Font(bold=True, size=14)
    s["A2"] = _xl(f"Input file: {source_name}")
    s["A3"] = (f"str_term_validator.py v{__version__} - dtSearch syntax + "
               "Search Terms Report constraints (each STR line runs as an "
               "individual dtSearch query)")
    s["A5"], s["B5"] = "Status", "Terms"
    s["A5"].font = s["B5"].font = Font(bold=True)
    row = 6
    counts = {st: 0 for st in STATUS_ORDER}
    for r in results:
        counts[r.status] += 1
    for st in STATUS_ORDER:
        s.cell(row=row, column=1, value=st).fill = \
            PatternFill("solid", fgColor=STATUS_FILL[st])
        s.cell(row=row, column=2, value=counts[st])
        row += 1
    s.cell(row=row, column=1, value="TOTAL").font = Font(bold=True)
    s.cell(row=row, column=2, value=len(results)).font = Font(bold=True)

    row += 2
    s.cell(row=row, column=1, value="Revised-term self-check").font = \
        Font(bold=True)
    row += 1
    for label in ("CLEAN", "ADVISORY", "NEEDS REVIEW"):
        s.cell(row=row, column=1, value=label)
        s.cell(row=row, column=2,
               value=sum(1 for r in results if r.residual.startswith(label)))
        row += 1

    row += 1
    s.cell(row=row, column=1, value="Most frequent issues").font = \
        Font(bold=True)
    row += 1
    freq = {}
    for r in results:
        for code in dict.fromkeys(i.code for i in r.issues):
            freq[code] = freq.get(code, 0) + 1
    for code, n in sorted(freq.items(), key=lambda kv: -kv[1])[:15]:
        s.cell(row=row, column=1, value=code)
        s.cell(row=row, column=2, value=n)
        row += 1

    row += 1
    s.cell(row=row, column=1, value=(
        "Caveats: noise-word list and alphabet-file behavior are index-"
        "specific defaults - confirm customized indexes with your Relativity "
        "administrator. Findings labeled ASSUMPTION APPLIED must be confirmed "
        "with the case team. Verify the hit-count impact of revisions in the "
        "workspace before relying on them."))
    s.cell(row=row, column=1).alignment = Alignment(wrap_text=True)
    s.column_dimensions["A"].width = 60
    s.column_dimensions["B"].width = 10

    wb.save(out_path)


def write_clean_txt(results, path):
    lines, seen = [], set()
    for r in results:
        t = r.revised.strip()
        if not t or r.residual.startswith("NEEDS REVIEW"):
            continue
        key = _collapse(t.lower())
        if key in seen:
            continue
        seen.add(key)
        lines.append(t)
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + ("\n" if lines else ""))
    return len(lines)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main(argv=None):
    global SPELLCHECK_ENABLED
    ap = argparse.ArgumentParser(
        description="Validate, auto-repair and self-verify Relativity Search "
                    "Terms Report terms (dtSearch syntax).")
    ap.add_argument("input", help="terms file: .xlsx, .csv or .txt")
    ap.add_argument("-o", "--output",
                    help="Excel report path (default: <input>_validation.xlsx)")
    ap.add_argument("--clean",
                    help="also write an STR-ready .txt of revised terms "
                         "(NEEDS REVIEW terms excluded, duplicates removed)")
    ap.add_argument("--no-spellcheck", action="store_true",
                    help="skip the optional spelling suggestions")
    ap.add_argument("--version", action="version",
                    version=f"%(prog)s {__version__}")
    args = ap.parse_args(argv)
    SPELLCHECK_ENABLED = not args.no_spellcheck

    if not os.path.exists(args.input):
        raise SystemExit(f"Input file not found: {args.input}")
    terms = read_terms(args.input)
    if not terms:
        raise SystemExit("No terms found in the input file.")

    results = validate_all(terms)

    out = args.output or os.path.splitext(args.input)[0] + "_validation.xlsx"
    write_excel(results, out, os.path.basename(args.input))

    counts = {}
    for r in results:
        counts[r.status] = counts.get(r.status, 0) + 1
    print(f"Validated {len(results)} term(s) from {args.input}")
    for st in STATUS_ORDER:
        if counts.get(st):
            print(f"  {st:8}: {counts[st]}")
    print(f"Report written to: {out}")
    rc = {"CLEAN": 0, "ADVISORY": 0, "NEEDS REVIEW": 0}
    for r in results:
        for k in rc:
            if r.residual.startswith(k):
                rc[k] += 1
                break
    print("Revised-term self-check: "
          f"{rc['CLEAN']} clean, {rc['ADVISORY']} advisory, "
          f"{rc['NEEDS REVIEW']} need review")
    if SPELLCHECK_ENABLED:
        import importlib.util
        if importlib.util.find_spec("spellchecker") is None:
            print("(spelling suggestions disabled - pip install "
                  "pyspellchecker to enable)")
    if args.clean:
        n = write_clean_txt(results, args.clean)
        print(f"Clean term list ({n} terms) written to: {args.clean}")


if __name__ == "__main__":
    main()
